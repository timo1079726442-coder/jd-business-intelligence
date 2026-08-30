"""Add validation dropdowns to configurable fields in config.xlsx."""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
path = ROOT / "config" / "config.xlsx"
wb = load_workbook(path)

boolean_values = '"是,否"'
enabled_values = '"启用,不启用"'
granularity_values = '"日,月,区间,默认"'
status_values = '"启用,停用"'

for ws in wb.worksheets:
    # Remove only validations created by this tool, keep unrelated validations.
    for dv in list(ws.data_validations.dataValidation):
        if dv.promptTitle == "配置选项":
            ws.data_validations.dataValidation.remove(dv)
    headers = {}
    for cell in ws[1]:
        if cell.value is not None:
            headers[str(cell.value).strip()] = cell.column
    for header, formula in (
        ("启用", boolean_values), ("是否启用", boolean_values),
        ("支持区间", boolean_values), ("默认粒度", granularity_values),
        ("粒度", granularity_values), ("状态", status_values),
    ):
        col = headers.get(header)
        if not col:
            continue
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.promptTitle = "配置选项"; dv.prompt = "请选择下拉选项，不要手工输入。"
        dv.errorTitle = "配置值无效"; dv.error = "请从下拉列表中选择有效值。"; dv.errorStyle = "stop"
        dv.showErrorMessage = True; dv.showInputMessage = True
        dv.add(f"{ws.cell(2, col).column_letter}2:{ws.cell(ws.max_row, col).column_letter}{max(ws.max_row, 500)}")
        ws.add_data_validation(dv)

wb.save(path)
# Re-open to verify validators survived serialization.
check = load_workbook(path, read_only=False)
validation_count = sum(len(ws.data_validations.dataValidation) for ws in check.worksheets)
check.close()
print(f"updated={path}")
print(f"validations={validation_count}")
