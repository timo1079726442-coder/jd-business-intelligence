import json, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from metrics.service import MetricService
p=Path('output/MIYO箱包旗舰店/京东MIYO数据库.xlsx')
wbv=load_workbook(p,read_only=True,data_only=True); wbf=load_workbook(p,read_only=True,data_only=False)
name=next(n for n in wbv.sheetnames if '全店' in n and wbv[n].max_row<=30 and wbv[n].max_column<=10)
v=wbv[name]; f=wbf[name]
legacy={"gmv":v['B4'].value,"refund_amount":v['B5'].value,"net_gmv":v['B6'].value,"refund_rate":v['B7'].value,"order_count":v['B9'].value,"refund_order_count":v['B10'].value,"return_rate":v['B11'].value,"ad_cost":v['B12'].value,"roi":v['B13'].value,"estimated_profit":v['B14'].value}
formulas={cell:f[cell].value for cell in ('B4','B5','B6','B7','B9','B10','B11','B12','B13','B14')}
api=MetricService().get_shop_overview('MIYO','2026-08-13','2026-08-19')
current={k:x['value'] for k,x in api['metrics'].items()}
print(json.dumps({'sheet':name,'legacy':legacy,'formulas':formulas,'api':current,'diff':{k:current.get(k)-legacy.get(k) if isinstance(current.get(k),(int,float)) and isinstance(legacy.get(k),(int,float)) else None for k in current}},ensure_ascii=False,indent=2,default=str))
