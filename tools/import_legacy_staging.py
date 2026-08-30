from pathlib import Path
import configparser, hashlib, json, re
from datetime import datetime
from openpyxl import load_workbook
import mysql.connector

root=Path(__file__).resolve().parents[1]
cfg=configparser.ConfigParser(); cfg.read(root/'config/mysql.local.ini',encoding='utf-8'); c=cfg['mysql']
db=mysql.connector.connect(host=c.get('host'),port=c.getint('port'),user=c.get('user'),password=c.get('password'),database=c.get('database'),charset=c.get('charset','utf8mb4'))
cur=db.cursor()
cur.execute('''CREATE TABLE IF NOT EXISTS legacy_excel_rows (
 row_id BIGINT AUTO_INCREMENT PRIMARY KEY, source_file VARCHAR(128) NOT NULL,
 source_shop VARCHAR(64) NOT NULL, source_sheet VARCHAR(128) NOT NULL,
 source_row INT NOT NULL, inferred_date DATE NULL, row_hash CHAR(64) NOT NULL,
 payload LONGTEXT NOT NULL, imported_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
 UNIQUE KEY uq_legacy_hash(row_hash), KEY ix_legacy_date(inferred_date), KEY ix_legacy_shop(source_shop)
) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci''')
db.commit()
keep={'订单列表','售后','推广自定义报表','快车订单','全站','全站-全店','全站订单','搜索','推荐','购物车'}
date_re=re.compile(r'(20\d{2})[\-/年](\d{1,2})[\-/月](\d{1,2})')
total=0
for path in sorted((root/'output'/'总表').glob('*.xlsx')):
 if not ('MIYO' in path.name or 'OTA' in path.name): continue
 shop='MIYO' if 'MIYO' in path.name else 'OTA'
 wb=load_workbook(path,read_only=True,data_only=True)
 for ws in wb.worksheets:
  if ws.title not in keep: continue
  batch=[]
  for n,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
   vals=[None if v is None else str(v) for v in row]
   text=' | '.join(v for v in vals if v)
   m=date_re.search(text)
   if not m: continue
   try: dt=datetime(int(m.group(1)),int(m.group(2)),int(m.group(3))).date()
   except ValueError: continue
   if dt >= datetime(2026,7,22).date(): continue
   payload=json.dumps(vals,ensure_ascii=False,separators=(',',':'))
   h=hashlib.sha256((path.name+'|'+ws.title+'|'+str(n)+'|'+payload).encode('utf-8')).hexdigest()
   batch.append((path.name,shop,ws.title,n,dt,h,payload))
   if len(batch)>=500:
    cur.executemany('INSERT IGNORE INTO legacy_excel_rows(source_file,source_shop,source_sheet,source_row,inferred_date,row_hash,payload) VALUES (%s,%s,%s,%s,%s,%s,%s)',batch); db.commit(); total+=len(batch); batch.clear()
  if batch:
   cur.executemany('INSERT IGNORE INTO legacy_excel_rows(source_file,source_shop,source_sheet,source_row,inferred_date,row_hash,payload) VALUES (%s,%s,%s,%s,%s,%s,%s)',batch); db.commit(); total+=len(batch)
  print(path.name,ws.title,'done')
 wb.close()
cur.execute('SELECT COUNT(*),MIN(inferred_date),MAX(inferred_date) FROM legacy_excel_rows')
print('RESULT',cur.fetchone(),'processed',total)
cur.close(); db.close()
