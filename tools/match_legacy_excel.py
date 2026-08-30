from pathlib import Path
import json, re, sqlite3
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
db=sqlite3.connect(f"file:{(root/'data/jd_report.db').as_posix()}?mode=ro",uri=True)
tables={}
for (t,) in db.execute("select name from sqlite_master where type='table' and name not like 'sqlite_%'"):
    tables[t]=[r[1] for r in db.execute(f'pragma table_info("{t}")')]
db.close()
files=[p for p in (root/'output'/'总表').glob('*.xlsx') if 'MIYO' in p.name or 'OTA' in p.name]
result=[]
for path in sorted(files):
    wb=load_workbook(path,read_only=True,data_only=True)
    book={'file':path.name,'sheets':[]}
    for ws in wb.worksheets:
        row=next(ws.iter_rows(min_row=1,max_row=1,values_only=True),())
        headers=[str(v).strip() for v in row if v is not None and str(v).strip()]
        norm=lambda x: re.sub(r'[^0-9a-zA-Z\u4e00-\u9fff]','',x).lower()
        hset={norm(x) for x in headers}
        scores=[]
        for t,cols in tables.items():
            cset={norm(x) for x in cols}
            scores.append((len(hset & cset)/max(1,len(hset|cset)),t))
        scores.sort(reverse=True)
        book['sheets'].append({'name':ws.title,'rows':ws.max_row,'cols':ws.max_column,'headers':headers[:20],'matches':scores[:3]})
    wb.close(); result.append(book)
out=root/'data'/'legacy_excel_matches.json'; out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
for b in result:
 print('\n'+b['file'])
 for s in b['sheets']:
  print(s['name'],s['rows'],s['matches'][0])
print('report',out)
