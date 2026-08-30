from pathlib import Path
from openpyxl import load_workbook
root = Path(__file__).resolve().parents[1]
folder = root / 'output' / 'MIYO箱包旗舰店' / 'MIYO商品主图对应ID'
print('folder_exists=', folder.exists())
for p in sorted(folder.iterdir()) if folder.exists() else []:
    print(p.name, p.suffix, p.stat().st_size)
for p in folder.glob('*.xlsx'):
    wb = load_workbook(p, read_only=True, data_only=True)
    print('FILE', p.name, 'SHEETS', wb.sheetnames)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
        print(ws.title, rows)
    wb.close()
