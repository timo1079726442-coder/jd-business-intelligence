"""建立 SKU→SPU→主图商品维表（当前先处理 MIYO）。"""
import configparser
import shutil
from pathlib import Path
from openpyxl import load_workbook
import mysql.connector

ROOT = Path(__file__).resolve().parents[1]
source_dir = ROOT / 'output' / 'MIYO箱包旗舰店' / 'MIYO商品主图对应ID'
mapping_file = source_dir / '商品ID货号.xlsx'
image_dir = ROOT / 'dashboard' / 'static' / 'products' / 'MIYO'
image_dir.mkdir(parents=True, exist_ok=True)

wb = load_workbook(mapping_file, read_only=True, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(min_row=2, values_only=True))
wb.close()
items = {}
for spu, sku, color, size in rows:
    if not spu or not sku:
        continue
    spu, sku = str(spu).split('.')[0], str(sku).split('.')[0]
    image_name = f'{spu}.png'
    image_src = source_dir / image_name
    image_path = f'/static/products/MIYO/{image_name}' if image_src.exists() else ''
    items[(spu, sku)] = (spu, sku, str(color or ''), str(size or ''), image_path, '已匹配' if image_path else '待补主图')
for spu in {x[0] for x in items.values()}:
    src = source_dir / f'{spu}.png'
    if src.exists():
        shutil.copy2(src, image_dir / src.name)

cfg = configparser.ConfigParser(); cfg.read(ROOT / 'config/mysql.local.ini', encoding='utf-8'); c = cfg['mysql']
db = mysql.connector.connect(host=c.get('host'), port=c.getint('port'), user=c.get('user'), password=c.get('password'), database=c.get('database'), charset=c.get('charset','utf8mb4'))
cur = db.cursor()
cur.execute('''CREATE TABLE IF NOT EXISTS dim_product_mapping (
 id BIGINT NOT NULL AUTO_INCREMENT PRIMARY KEY,
 shop_code VARCHAR(32) NOT NULL, spu_id VARCHAR(64) NOT NULL, sku_id VARCHAR(64) NOT NULL,
 sku_color VARCHAR(128), sku_size VARCHAR(128), image_path VARCHAR(255), mapping_status VARCHAR(32) NOT NULL,
 UNIQUE KEY uk_shop_sku (shop_code, sku_id), KEY ix_shop_spu (shop_code, spu_id)
) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci''')
cur.execute("DELETE FROM dim_product_mapping WHERE shop_code='MIYO'")
cur.executemany('''INSERT INTO dim_product_mapping (shop_code,spu_id,sku_id,sku_color,sku_size,image_path,mapping_status)
 VALUES (%s,%s,%s,%s,%s,%s,%s)''', [('MIYO',) + x for x in items.values()])
db.commit(); print(f'MIYO_MAPPING_OK rows={len(items)} images={sum(1 for x in items.values() if x[4])}')
cur.close(); db.close()
