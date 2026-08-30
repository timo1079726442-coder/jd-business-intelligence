from pathlib import Path
import json,re,sqlite3
from datetime import datetime
from collections import defaultdict
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
db=sqlite3.connect(f"file:{(root/'data/jd_report.db').as_posix()}?mode=ro",uri=True)
tables={}
for (t,) in db.execute("select name from sqlite_master where type='table' and name not like 'sqlite_%'"):
 tables[t]=[r[1] for r in db.execute(f'pragma table_info("{t}")')]
db.close()
norm=lambda x: re.sub(r'[^0-9a-zA-Z\u4e00-\u9fff]','',str(x or '')).lower()
date_re=re.compile(r'(20\d{2})[\-/年](\d{1,2})[\-/月](\d{1,2})')
mapping={'订单列表':'biz_jm_order_full','售后':'biz_jm_after_sale_full','推广自定义报表':'biz_jzt_kuaiche','快车订单':'biz_jzt_kuaiche_order_effect','全站':'biz_jzt_quanzhan_campaign','全站-全店':'biz_jzt_quanzhan_campaign_all_store','全站订单':'biz_jzt_quanzhan_effect','搜索':'biz_traffic_search','推荐':'biz_traffic_recommend','购物车':'biz_traffic_cart'}
stats=defaultdict(lambda:{'rows':0,'min':None,'max':None,'matched_columns':0})
for path in sorted((root/'output'/'总表').glob('*.xlsx')):
 if 'MIYO' not in path.name and 'OTA' not in path.name: continue
 wb=load_workbook(path,read_only=True,data_only=True)
 for ws in wb.worksheets:
  target=mapping.get(ws.title)
  if not target: continue
  headers=next(ws.iter_rows(min_row=1,max_row=1,values_only=True),())
  cset={norm(x) for x in tables[target]}; matched=sum(norm(x) in cset for x in headers if x is not None)
  key=(path.stem,ws.title,target); s=stats[key]; s['matched_columns']=matched
  for row in ws.iter_rows(min_row=2,values_only=True):
   text=' | '.join(str(v) for v in row if v is not None); m=date_re.search(text)
   if not m: continue
   try: dt=datetime(int(m.group(1)),int(m.group(2)),int(m.group(3))).date()
   except ValueError: continue
   if dt>=datetime(2026,7,22).date(): continue
   s['rows']+=1; s['min']=str(dt) if s['min'] is None or str(dt)<s['min'] else s['min']; s['max']=str(dt) if s['max'] is None or str(dt)>s['max'] else s['max']
 wb.close()
out={"status":"PREVIEW_ONLY","items":[dict(file=k[0],sheet=k[1],target=k[2],**v) for k,v in stats.items()]}
(root/'data'/'legacy_merge_preview.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
for x in out['items']: print(x)
print('PREVIEW_ROWS',sum(x['rows'] for x in out['items']))
