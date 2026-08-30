from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
p = root / 'output' / 'MIYO箱包旗舰店' / '京东MIYO数据库.xlsx'
wb = load_workbook(p, read_only=True, data_only=False)
print('SHEETS', wb.sheetnames)
for ws in wb.worksheets:
    if ws.title not in ('全店复盘', '商品复盘', '单品分日'):
        continue
    print('\nSHEET', ws.title, 'DIM', ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 16), values_only=False):
        vals = []
        for cell in row[:min(ws.max_column, 12)]:
            if cell.value is not None:
                vals.append(f'{cell.coordinate}={cell.value!r}')
        if vals:
            print(' | '.join(vals))
wb.close()
