from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
p = root / 'output' / 'MIYO箱包旗舰店' / '京东MIYO数据库.xlsx'
wb = load_workbook(p, read_only=True, data_only=False)
ws = wb['订单列表']
for i in range(75, 84):
    print(i + 1, ws.cell(1, i + 1).value)
wb.close()
