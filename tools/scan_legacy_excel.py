from pathlib import Path
import json
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
folder=root/'output'/'总表'
result=[]
paths=[p for p in folder.glob('*.xlsx') if 'MIYO' in p.name or 'OTA' in p.name]
for path in sorted(paths):
    wb=load_workbook(path, read_only=True, data_only=True)
    book={'file':path.name,'sheets':[]}
    for ws in wb.worksheets:
        rows=ws.iter_rows(min_row=1,max_row=3,values_only=True)
        sample=[]
        for row in rows:
            sample.append([str(v)[:120] if v is not None else None for v in row[:25]])
        book['sheets'].append({'name':ws.title,'rows':ws.max_row,'cols':ws.max_column,'sample':sample})
    wb.close()
    result.append(book)
out=root/'data'/'legacy_excel_preflight.json'
out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print('files', [x['file'] for x in result], 'report', out)
