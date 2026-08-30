from pathlib import Path
import configparser, json
from openpyxl import load_workbook
import mysql.connector
root=Path(__file__).resolve().parents[1]
cfg=configparser.ConfigParser(); cfg.read(root/'config/mysql.local.ini',encoding='utf-8'); c=cfg['mysql']
db=mysql.connector.connect(host=c.get('host'),port=c.getint('port'),user=c.get('user'),password=c.get('password'),database=c['database'],charset=c.get('charset','utf8mb4'))
cur=db.cursor(); cur.execute('''CREATE TABLE IF NOT EXISTS legacy_excel_sheet_meta(source_file VARCHAR(128),source_sheet VARCHAR(128),headers_json LONGTEXT,PRIMARY KEY(source_file,source_sheet)) CHARACTER SET utf8mb4''')
for path in sorted((root/'output'/'总表').glob('*.xlsx')):
 if 'MIYO' not in path.name and 'OTA' not in path.name: continue
 wb=load_workbook(path,read_only=True,data_only=True)
 for ws in wb.worksheets:
  row=next(ws.iter_rows(min_row=1,max_row=1,values_only=True),())
  headers=[None if v is None else str(v) for v in row]
  cur.execute('REPLACE INTO legacy_excel_sheet_meta VALUES (%s,%s,%s)',(path.name,ws.title,json.dumps(headers,ensure_ascii=False)))
 wb.close()
db.commit(); cur.execute('SELECT COUNT(*) FROM legacy_excel_sheet_meta'); print('SHEET_META',cur.fetchone()[0])
cur.close();db.close()
