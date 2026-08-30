"""Export a human-readable mapping of MySQL tables to JD business sections."""
import configparser
from pathlib import Path
import mysql.connector
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
cfg = configparser.ConfigParser(); cfg.read(ROOT / "config" / "mysql.local.ini", encoding="utf-8")
db = cfg["mysql"]
mapping = [
    ("京麦", "订单", "std_biz_jm_order_full", "vw_dashboard_orders_daily", "订单明细；成交金额、退款金额、订单数", "单日/区间补数；京麦订单近30天窗口"),
    ("京麦", "售后", "std_biz_jm_after_sale_full", "vw_dashboard_shop_daily", "售后/退款明细；退款金额、退货订单", "单日/区间补数；京麦售后近30天窗口"),
    ("京准通", "快车推广", "std_biz_jzt_kuaiche", "vw_dashboard_promotion_daily", "快车推广消耗、曝光、点击、转化等", "可按缺失区间导出；空表表示无投放"),
    ("京准通", "快车订单效果", "std_biz_jzt_kuaiche_order_effect", "vw_dashboard_promotion_daily", "快车带来的订单效果", "按平台支持的区间补数"),
    ("京准通", "全站推广计划", "std_biz_jzt_quanzhan_campaign", "vw_dashboard_promotion_daily", "全站计划/推广成本", "按平台支持的区间补数；空表表示无投放"),
    ("京准通", "全站全店", "std_biz_jzt_quanzhan_campaign_all_store", "vw_dashboard_promotion_daily", "全站全店维度推广数据", "按平台支持的区间补数；空表表示无投放"),
    ("京准通", "全站订单效果", "std_biz_jzt_quanzhan_effect", "vw_dashboard_promotion_daily", "全站推广订单效果", "按平台支持的区间补数"),
    ("商智", "搜索流量", "std_biz_traffic_search", "vw_dashboard_shop_daily", "搜索来源流量、访客、点击等", "可区间则区间；仅支持单日则按日"),
    ("商智", "推荐流量", "std_biz_traffic_recommend", "vw_dashboard_shop_daily", "推荐来源流量", "可区间则区间；仅支持单日则按日"),
    ("商智", "购物车流量", "std_biz_traffic_cart", "vw_dashboard_shop_daily", "购物车来源流量", "可区间则区间；仅支持单日则按日"),
    ("统一层", "商品映射", "dim_product_mapping", "—", "SKU→SPU→商品主图；缺失标记待补映射", "维护商品ID货号.xlsx及静态主图目录"),
    ("统一层", "历史 Excel 原始层", "legacy_excel_rows", "—", "旧版 Excel 补充历史记录", "仅作历史补充，不覆盖新标准表"),
    ("统一层", "历史字段元数据", "legacy_excel_sheet_meta", "—", "旧 Excel 工作表字段/表头元数据", "用于字段映射审查"),
    ("统一层", "统一历史业务层", "unified_legacy_rows", "vw_unified_business", "旧版数据统一查询层", "兼容历史看板与追溯"),
    ("运行记录", "迁移运行记录", "mysql_migration_runs", "—", "SQLite→MySQL迁移批次、状态、行数", "用于审计与故障追踪"),
]

conn = mysql.connector.connect(host=db.get("host"), port=db.getint("port"), user=db.get("user"), password=db.get("password"), database=db.get("database"))
cur = conn.cursor()
rows = []
for section, name, table, view, purpose, rule in mapping:
    cur.execute("SELECT COUNT(*) FROM information_schema.tables WHERE table_schema=%s AND table_name=%s", (db.get("database"), table))
    exists = cur.fetchone()[0] == 1
    count = ""; date_range = ""
    if exists:
        try:
            cur.execute(f"SELECT COUNT(*) FROM `{table}`"); count = cur.fetchone()[0]
        except Exception:
            count = "读取失败"
    rows.append([section, name, table, view, "是" if exists else "否", count, purpose, rule])
cur.close(); conn.close()

wb = Workbook(); ws = wb.active; ws.title = "业务表映射"
headers = ["一级板块", "数据板块", "MySQL表", "关联视图", "当前存在", "记录数", "数据用途", "采集/补数规则"]
ws.append(headers)
for row in rows: ws.append(row)
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="173B63"); cell.alignment = Alignment(horizontal="center")
widths = [12, 18, 38, 32, 10, 14, 48, 54]
for i, width in enumerate(widths, 1): ws.column_dimensions[chr(64+i)].width = width
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)

note = wb.create_sheet("看板口径")
note.append(["指标", "口径/来源"])
notes = [
    ("复盘期", "用户选择的起止日期，包含首尾两天"),
    ("对比期", "与复盘期等长的紧邻前一周期；例如近7天对比上一个连续7天"),
    ("成交金额", "订单表订单数=1的订单拆分金额合计"),
    ("退款金额", "订单表订单数=1的订单号汇总退款金额合计"),
    ("净成交额", "成交金额－退款金额"),
    ("总订单数", "订单数=1、出库状态=已出库、成交金额>10的记录数"),
    ("退货订单数", "订单数=1、出库状态=已出库、退款金额>10的记录数"),
    ("退款率/退货率", "退款金额/成交金额；退货订单数/总订单数"),
    ("推广花费", "快车花费 + 全站计划投放成本 + 全站全店投放成本"),
    ("空推广表", "平台返回空表按无推广处理，不人为补零之外的示例数据"),
]
for row in notes: note.append(row)
note.freeze_panes = "A2"; note.column_dimensions["A"].width = 20; note.column_dimensions["B"].width = 100
for cell in note[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="173B63")
for row in note.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)

out = ROOT / "output" / "数据库表业务映射.xlsx"; out.parent.mkdir(parents=True, exist_ok=True); wb.save(out)
check = load_workbook(out, read_only=True); assert check.sheetnames == ["业务表映射", "看板口径"]; assert check["业务表映射"].max_row == len(mapping) + 1; check.close()
print(out)
