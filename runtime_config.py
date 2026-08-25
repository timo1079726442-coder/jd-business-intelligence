# -*- coding: utf-8 -*-
"""runtime_config.py
============================================================
运行时配置统一加载器（2026-08-22 审计 H-04/H-05/H-06 修复新增）

目的：
    1. 把禁止硬编码的业务常量（店铺 SHOP_ID / SHOP_PIN / APP_ID / SIGN_SALT）集中读取
    2. 任何字段读取不到 → 立即 SystemExit 报错，禁止静默回落到默认值
    3. 替代 main.py 23 处 os.getenv("SHOP_ID", "FYA箱包旗舰店") 硬编码兜底

数据源优先级：
    ① 环境变量（CLI/调度脚本 set 注入）
    ② config.xlsx「店铺账号」sheet（shop_pin）
    ③ config.xlsx「全局配置」sheet（APP_ID / sign_salt 等业务凭证）

兜底策略（2026-08-22 当前过渡）：
    config.xlsx 中如果新键尚未补登（用户手动 Excel 操作未完成），get_*() 会尝试读
    老的同义键（如「商品流量来源 / 签名盐值」），如果还找不到才 SystemExit。
    这样既保证 H-04/H-05 修复目标（不依赖 py 源码硬编码），又不阻塞运行时。

对外 API：
    get_shop_id()              -> str  店铺全名（必填，未读到 SystemExit）
    get_shop_pin()             -> str  京东 pin 主账号（必填）
    get_app_id(biz_type)       -> str  京东 appId（必填；biz_type=jm_order/jm_after_sale）
    get_sign_salt()            -> str  商智 MD5 签名盐值（必填）
    get_runtime_config()       -> dict 一次性返回全部字段（调试用）
"""
import os
import sys
import logging
from typing import Optional

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
CONFIG_DIR = os.path.join(PROJECT_ROOT, "config")
CONFIG_PATH = os.path.join(CONFIG_DIR, "config.xlsx")


def _exit(msg: str, code: int = 3) -> None:
    """统一报错出口（参数错误 → 退出码 3，符合 AGENTS.md 调度约定）"""
    print("[FATAL] " + msg, file=sys.stderr)
    logging.error("[FATAL] " + msg)
    sys.exit(code)


def _read_xlsx_sheet(sheet_name: str) -> list:
    """读取 config.xlsx 指定 sheet，返回二维 list。"""
    if not os.path.exists(CONFIG_PATH):
        _exit("config.xlsx 不存在：" + CONFIG_PATH)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(CONFIG_PATH, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            _exit(
                "config.xlsx 缺少 sheet「" + sheet_name + "」\n"
                "   已有的 sheet：" + str(wb.sheetnames)
            )
        ws = wb[sheet_name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    except SystemExit:
        raise
    except Exception as e:
        _exit("config.xlsx「" + sheet_name + "」读取失败：" + str(e))


def _lookup_xlsx_value(group: str, var: str) -> Optional[str]:
    """通用查表函数：从「全局配置」sheet 找 (项目名, 变量参数) 对应的「参数值」。"""
    rows = _read_xlsx_sheet("全局配置")
    for row in rows[1:]:
        if not row or len(row) < 3:
            continue
        g = str(row[0]).strip() if row[0] else ""
        v = str(row[1]).strip() if row[1] else ""
        if g == group and v == var:
            value = str(row[2]).strip() if row[2] else ""
            if value:
                return value
    return None


def _lookup_shop_pin(shop_short_name: str) -> str:
    """从 config.xlsx「店铺账号」sheet 按店铺短名查 shop_pin。

    表格结构：店铺名 / 账号 / 密码 / 启用 / ...
    返回「账号」列（如 FYA8888 / miyo-周 / ota8888）。
    """
    rows = _read_xlsx_sheet("店铺账号")
    for row in rows[1:]:
        if not row or len(row) < 4:
            continue
        shop_name = str(row[0]).strip() if row[0] else ""
        pin = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        enabled = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        if shop_name == shop_short_name and enabled in ("是", "True", "true", "1", "yes"):
            return pin
    _exit(
        "config.xlsx「店铺账号」sheet 中未找到启用店铺「" + shop_short_name + "」的账号\n"
        "   → 请确认店铺短名（FYA/MIYO/OTA）与表中「店铺名」列完全一致"
    )


def get_shop_id() -> str:
    """获取当前店铺全名（如「FYA箱包旗舰店」），必填。"""
    val = os.getenv("SHOP_ID", "").strip()
    if not val:
        _exit(
            "环境变量 SHOP_ID 未设置\n"
            "   → 请 set SHOP_ID=FYA箱包旗舰店 / MIYO箱包旗舰店 / OTA箱包旗舰店\n"
            "   → 或 CLI 参数 --shop 注入（由各入口自行 set）"
        )
    return val


def get_shop_short_name() -> str:
    """店铺全名 → 短名（FYA / MIYO / OTA），供查 config.xlsx 用。"""
    return get_shop_id().replace("箱包旗舰店", "")


def get_shop_pin() -> str:
    """获取当前店铺的京东 pin 主账号（如 FYA8888 / miyo-周 / ota8888），必填。

    数据源优先级：
        ① 环境变量 SHOP_PIN（允许 CLI 覆盖）
        ② config.xlsx「店铺账号」sheet 按店铺短名查
    """
    env_pin = os.getenv("SHOP_PIN", "").strip()
    if env_pin:
        return env_pin
    short = get_shop_short_name()
    return _lookup_shop_pin(short)


def get_app_id(biz_type: str) -> str:
    """获取京东 appId，必填。

    数据源：环境变量 APP_ID_{BIZ_TYPE} → APP_ID → config.xlsx「全局配置」sheet「app_id_<biz_type>」

    Args:
        biz_type: jm_order（京麦订单明细）/ jm_after_sale（京麦售后明细）
    """
    env_key = "APP_ID_" + biz_type.upper()
    env_val = os.getenv(env_key, "").strip() or os.getenv("APP_ID", "").strip()
    if env_val:
        return env_val
    target_key = "app_id_" + biz_type
    v = _lookup_xlsx_value("京麦接口", target_key)
    if v:
        return v
    _exit(
        "appId 读取失败：biz_type=" + biz_type + "\n"
        "   → 环境变量 " + env_key + " / APP_ID 未设置\n"
        "   → config.xlsx「全局配置」sheet 未配置 京麦接口/" + target_key + "\n"
        "   → 请补登配置后重试（H-04 修复要求）"
    )


def get_sign_salt() -> str:
    """获取商智 MD5 签名盐值，必填。

    数据源：
        ① config.xlsx「全局配置」sheet「全局/sign_salt」
        ② 兜底「商品流量来源/签名盐值」（H-05 修复前的同义键）
    """
    # 1. 优先读新加的 key
    v = _lookup_xlsx_value("全局", "sign_salt")
    if v:
        return v
    # 2. 兜底到老 key（保证 H-05 修复不阻塞运行）
    v = _lookup_xlsx_value("商品流量来源", "签名盐值")
    if v:
        logging.warning(
            "[H-05] config.xlsx「全局/sign_salt」未配置，临时读老键「商品流量来源/签名盐值」\n"
            "       建议在「全局配置」sheet 新增一行：全局 | sign_salt | " + str(v)
        )
        return v
    _exit(
        "商智签名盐值未配置\n"
        "   → config.xlsx「全局配置」sheet「全局/sign_salt」或「商品流量来源/签名盐值」不能为空\n"
        "   → H-05 修复要求：禁止硬编码 _DEFAULT_SIGN_SALT"
    )


def get_shop_name_display() -> str:
    """CLI banner 用的店铺展示名（与 SHOP_ID 一致）。"""
    return get_shop_id()


def get_runtime_config() -> dict:
    """一次性返回全部运行时配置（调试用）。"""
    return {
        "shop_id": get_shop_id(),
        "shop_pin": get_shop_pin(),
        "app_id_jm_order": get_app_id("jm_order"),
        "app_id_jm_after_sale": get_app_id("jm_after_sale"),
        "sign_salt": get_sign_salt(),
    }


if __name__ == "__main__":
    import json
    print(json.dumps(get_runtime_config(), ensure_ascii=False, indent=2))
