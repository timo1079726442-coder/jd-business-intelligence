# -*- coding: utf-8 -*-
"""多店铺数据导出手动触发入口（项目22，2026-08-17）

用户决策 2026-08-17：
    - **手动触发**：本脚本是手动工具，不接影刀定时调度。你想跑就执行一次。
    - **不要凌晨自动跑**：鉴权过期时人工判断重抓更稳。
    - **3 个店一次跑完**：传 --shop 跑单店；不传则遍历 3 个店。

职责：
    1. 每个店按"近 30 天"拉一次：京麦订单 + 京麦售后 + 京准通全 6 业务
    2. 商智（搜索/购物车/推荐 + 4 个其他业务）走缺日补齐：
       - 先扫 output/{店名}/{业务}/{date}/ 是否存在 → 缺哪些日期
       - 只补缺的那些日期（避免重复请求触发风控）
    3. IMAP 收件统一用一份 config/imap_config.ini（QQ:1079726442@qq.com）
    4. 按退出码约定返回：
       0 = 全部成功
       1 = 部分业务失败
       2 = 鉴权过期（你需手动触发 RPA 重抓 cookie/h5st）
       3 = 参数错误 / 系统错误

调用示例：
    # 全部3 个店全跑：
    python rpa_run.py
    # 单店跑：
    python rpa_run.py --shop "FYA箱包旗舰店"
    # Dry-run（仅预览，不实际下载）：
    python rpa_run.py --dry-run
"""
import os
import sys
import time
import argparse
import traceback
from datetime import datetime, timedelta

# Phase 2.3（2026-08-20）：店铺/业务清单从 config.xlsx「店铺清单」「业务清单」sheet 读取
# 替代原硬编码 KNOWN_SHOPS / JM_JZT_BIZ_KEYS / 商智_BIZ_KEYS / biz_to_dirname（AGENTS.md 第3条）
from biz_config_loader import list_shop_ids, list_biz_keys, get_output_dirname

# 必须在 import main 之前设环境变量（启用 AuthLoader 接管 cookie）
if "AUTH_LOADER" not in os.environ:
    os.environ["AUTH_LOADER"] = "1"
if "AUTH_RPA_CLI" not in os.environ:
    os.environ["AUTH_RPA_CLI"] = "0"  # 手动触发模式：不调 RPA
# Cookie 过期检查跳过（AGENTS.md 用户决策：由接口返回码判定）
os.environ["AUTH_SKIP_COOKIE_EXPIRE_CHECK"] = "1"
# h5st 过期检查跳过（2026-08-20：与 Cookie 同理，由接口返回码判定是否失效）
os.environ["AUTH_SKIP_H5ST_EXPIRE_CHECK"] = "1"

# 启动时 IMAP 健康检查（2026-08-21 项目23 新增）
# - 检查 config.xlsx「全局配置」sheet 的 IMAP 组 auth_code_expire_date
# - 距过期 < 14 天 WARN；< 7 天或已过期抛 RuntimeError（京麦订单明细会失败，提前拦）
# - 用户决策：3 个月更新一次授权码，到期前一周预警
from imap_config_loader import load_imap_config, check_imap_health as _imap_check
try:
    _imap_cfg = load_imap_config()
    _imap_check(_imap_cfg)
except (FileNotFoundError, ValueError) as e:
    # ini 缺失 → 静默跳过（项目内只有京麦订单明细需要 IMAP，未配置时跳过）
    print(f"[IMAP] 跳过健康检查：{e}")
except RuntimeError as e:
    # 授权码过期 → 阻塞整个调度（京麦订单明细会失败）
    print(f"\n[BLOCK] {e}\n")
    # 注意：用户决策是阻塞，但允许 dry-run 继续
    # （如果用户只是预览计划，应该放过；只有真正执行时才报错）
    # 这里先打印，由 main() 中根据 dry_run 决定是否 raise
    _IMAP_BLOCK_ERROR = e
else:
    _IMAP_BLOCK_ERROR = None

# ============================ 退出码 ============================
EXIT_SUCCESS = 0
EXIT_BIZ_FAIL = 1
EXIT_AUTH_EXPIRED = 2
EXIT_SYSTEM_ERROR = 3

# ============================ 业务清单（Phase 2.3 改为从 config.xlsx 读取）============================
# 替代原硬编码 KNOWN_SHOPS / JM_JZT_BIZ_KEYS / 商智_BIZ_KEYS（AGENTS.md 第3条禁止硬编码业务列表）
# 调用方式：
#   list_shop_ids()                       → 启用店铺清单（替代 KNOWN_SHOPS）
#   list_biz_keys(batch="jm_jzt")         → 京麦+京准通批次（替代 JM_JZT_BIZ_KEYS）
#   list_biz_keys(batch="sz")             → 商智批次（替代 商智_BIZ_KEYS）
#   get_output_dirname(biz_key)          → 输出目录名（替代 biz_to_dirname 字典）
# 数据源：config.xlsx「店铺清单」「业务清单」两个 sheet（由 create_config_sheets.py 生成）


def compute_last_30_days():
    """计算「今天-30天 ~ 今天-1天」日期区间（AGENTS.md 每日定时约定）

    返回:
        (start_date, end_date) - "YYYY-MM-DD" 格式字符串
    """
    today = datetime.now()
    end = today - timedelta(days=1)
    start = today - timedelta(days=30)
    return start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d")


def compute_last_n_days(n: int):
    """计算「今天-N天 ~ 今天-1天」日期区间（2026-08-20 新增，支持商智近7天等自定义窗口）

    参数:
        n - 天数（如 7 = 近7天）

    返回:
        (start_date, end_date) - "YYYY-MM-DD" 格式字符串
    """
    today = datetime.now()
    end = today - timedelta(days=1)
    start = today - timedelta(days=n)
    return start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d")


def _read_shop_interval_seconds(cli_value):
    """读取店间安全间隔秒数（2026-08-21 项目23 新增）。

    优先级：
        1. CLI 参数 --shop-interval（最优先）
        2. config.xlsx「全局配置」sheet 全局组 → shop_interval_seconds
        3. 兜底 60 秒

    参数:
        cli_value: 命令行 --shop-interval 值；None → 走 xlsx

    返回:
        int - 间隔秒数（>= 0）
    """
    if cli_value is not None:
        if cli_value < 0:
            print(f"[WARN] --shop-interval 不能为负数（{cli_value}），使用 0（不等待）")
            return 0
        return int(cli_value)

    # 从 config.xlsx 读取
    try:
        import openpyxl
        xlsx_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config", "config.xlsx")
        if not os.path.exists(xlsx_path):
            print("[WARN] config.xlsx 不存在，使用兜底店间间隔 60 秒")
            return 60
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        for sheet_name in ("全局配置", "全局", "config"):
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 3:
                        continue
                    group = str(row[0]).strip() if row[0] else ""
                    key = str(row[1]).strip() if row[1] else ""
                    val = row[2]
                    if group == "全局" and key == "shop_interval_seconds" and val is not None:
                        try:
                            sec = int(str(val).strip())
                            wb.close()
                            return max(0, sec)
                        except (ValueError, TypeError):
                            print(f"[WARN] shop_interval_seconds 值非法：{val}，使用 60 秒")
                            wb.close()
                            return 60
                break
        wb.close()
        # 未找到配置项
        return 60
    except Exception as e:
        print(f"[WARN] config.xlsx 读取失败（{e}），使用兜底 60 秒")
        return 60


def list_existing_dates(output_root: str, shop_id: str, biz_key: str) -> set:
    """扫描 output/{店名}/{业务名}/ 下文件名中的日期，返回已有日期集合（2026-08-21 改平铺后从文件名提取日期）

    平铺后的目录结构：
        output/{店名}/{业务名}/{业务名}_{date}.xlsx
        output/{店名}/{业务名}/{业务名}_{date}_{granularity}.xlsx
    本函数从 .xlsx 文件名中提取 YYYY-MM-DD 段，等价于原"扫描日期子目录"行为。

    参数:
        output_root - 项目根/output
        shop_id     - 店铺名
        biz_key - 业务 key（用来定位子目录名）

    返回:
        set - 已存在的日期集合 {"2026-07-18", "2026-07-19", ...}
    """
    # Phase 2.3：输出目录名从 config.xlsx「业务清单」sheet 读取（替代原硬编码 biz_to_dirname 字典）
    dirname = get_output_dirname(biz_key)
    if not dirname or dirname == biz_key:
        # 兜底：biz_key 未在 config.xlsx 注册时返回空集合（视为该业务无已有日期）
        return set()

    biz_dir = os.path.join(output_root, shop_id, dirname)
    if not os.path.isdir(biz_dir):
        return set()

    import re
    # YYYY-MM-DD 日期正则：匹配 2026-07-18 形式
    date_pattern = re.compile(r"(20\d{2}-\d{2}-\d{2})")

    existing = set()
    for entry in os.listdir(biz_dir):
        full = os.path.join(biz_dir, entry)
        # 平铺模式：扫描 .xlsx 文件，从文件名提取日期
        if os.path.isfile(full) and entry.lower().endswith(".xlsx"):
            m = date_pattern.search(entry)
            if m:
                # 校验日期合法性（如 2026-13-45 这种非法日期会被丢弃）
                date_str = m.group(1)
                try:
                    datetime.strptime(date_str, "%Y-%m-%d")
                    existing.add(date_str)
                except ValueError:
                    continue
        # 兼容旧数据：如果还存在日期子目录（2026-08-21 改造前的旧文件），也读取子目录名作为日期
        elif os.path.isdir(full) and len(entry) == 10 and entry[4] == "-":
            existing.add(entry)
    return existing


def compute_missing_dates(start_date: str, end_date: str, existing: set) -> list:
    """计算 [start_date, end_date] 区间里缺失的日期列表（按时间正序）

    返回:
        list - 缺失日期字符串列表 ["2026-08-15", "2026-08-16", ...]
    """
    start_dt = datetime.strptime(start_date, "%Y-%m-%d")
    end_dt = datetime.strptime(end_date, "%Y-%m-%d")
    missing = []
    cur = start_dt
    while cur <= end_dt:
        s = cur.strftime("%Y-%m-%d")
        if s not in existing:
            missing.append(s)
        cur += timedelta(days=1)
    return missing


def run_shop(shop_id: str, dry_run: bool = False, sz_days: int = 30, skip_jm: bool = False, only_jm: bool = False) -> int:
    """跑单个店的所有业务

    参数:
        shop_id  - 店铺名
        dry_run  - 仅预览不实际下载
        sz_days  - 商智业务日期窗口天数（默认30，用户可设7等）
        skip_jm  - 跳过京麦业务（避开频控窗口，只跑京准通+商智）
        only_jm  - 只跑京麦业务（与 skip_jm 互斥；跳过京准通+商智）

    返回:
        int - 该店退出码（0=全成功 / 1=部分失败 / 2=鉴权过期）
    """
    print(f"\n{'='*70}")
    print(f"[SHOP] {shop_id} 开始执行")
    print(f"{'='*70}")

    os.environ["SHOP_ID"] = shop_id  # 确保 AuthLoader 读到正确店
    # H-11 修复（2026-08-24 审计）：同时 set SHOP_PIN，避免 db_utils DELETE 错位
    # 背景：db_utils.upsert_df fallback 用 shop_id.replace("箱包旗舰店","")="FYA"，
    #       但 config 实际是 shop_pin="FYA8888"，DELETE 找不到 → 数据堆积
    try:
        from runtime_config import get_shop_pin
        os.environ["SHOP_PIN"] = get_shop_pin()
    except SystemExit as e:
        print(f"[WARN] 店铺 {shop_id} 的 SHOP_PIN 读取失败：{e}")
        print(f"[WARN] 继续跑业务，但 DB 入库可能错位（建议检查 config.xlsx「店铺账号」sheet）")

    # 延迟 import（环境变量设完后再 import main）
    try:
        from main import run_business
        from auth_loader import CookieExpiredError, H5stExpiredError
    except ImportError as e:
        print(f"[ERR] import main 失败：{e}")
        return EXIT_SYSTEM_ERROR

    # 京麦 + 京准通：近 30 天全量覆盖
    jm_jzt_start, jm_jzt_end = compute_last_30_days()
    print(f"[DATE] 京麦+京准通 近30天: {jm_jzt_start} ~ {jm_jzt_end}")

    # 商智：近 N 天（用户可指定，默认30，本次验证用7）
    sz_start, sz_end = compute_last_n_days(sz_days)
    print(f"[DATE] 商智 近{sz_days}天: {sz_start} ~ {sz_end}")

    overall_exit = EXIT_SUCCESS
    has_auth_expired = False

    # === 第 1 批：京麦 + 京准通（按"近 30 天"全量跑，覆盖原日期）===
    # Phase 2.3：业务清单从 config.xlsx「业务清单」sheet 读取（batch=jm_jzt）
    # 2026-08-20：支持 --skip-jm 跳过京麦业务（避开频控窗口）
    # 2026-08-21：支持 --only-jm 只跑京麦业务（与 skip_jm 互斥）
    if only_jm:
        batch_label = "京麦（--only-jm 模式，跳过京准通+商智）"
    elif skip_jm:
        batch_label = "京准通（跳过京麦）"
    else:
        batch_label = "京麦 + 京准通"
    print(f"\n--- [BATCH 1/2] {batch_label}：近 30 天全量覆盖 ---")
    for biz_key in list_biz_keys(batch="jm_jzt"):
        # --only-jm：跳过非京麦业务
        if only_jm and "京麦" not in biz_key:
            print(f"[SKIP-NON-JM] {biz_key}（--only-jm 模式，跳过非京麦业务）")
            continue
        # --skip-jm：跳过京麦业务
        if skip_jm and "京麦" in biz_key:
            print(f"[SKIP-JM] {biz_key}（跳过京麦，避开频控窗口）")
            continue
        if dry_run:
            print(f"[DRY-RUN] {biz_key} | {jm_jzt_start} ~ {jm_jzt_end}")
            continue
        print(f"\n>>> {biz_key} ({jm_jzt_start} ~ {jm_jzt_end})")
        try:
            run_business(biz_key, start_date=jm_jzt_start, end_date=jm_jzt_end)
            print(f"[OK] {biz_key}")
        except (CookieExpiredError, H5stExpiredError) as e:
            print(f"[AUTH_EXPIRED] {biz_key}: {e}")
            has_auth_expired = True
        except Exception as e:
            print(f"[BIZ_FAIL] {biz_key}: {e}")
            overall_exit = EXIT_BIZ_FAIL

    # === 第 2 批：商智业务（按"缺日补齐"，窗口近 N 天）===
    # Phase 2.3：业务清单从 config.xlsx「业务清单」sheet 读取（batch=sz）
    # 2026-08-20：商智窗口可自定义（--sz-days），默认30天
    # 2026-08-21：--only-jm 模式整体跳过商智批次
    if only_jm:
        print(f"\n--- [BATCH 2/2] 商智业务：跳过（--only-jm 模式只跑京麦）---")
    else:
        print(f"\n--- [BATCH 2/2] 商智业务：近{sz_days}天缺日补齐 ---")
        output_root = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")

        for biz_key in list_biz_keys(batch="sz"):
            existing = list_existing_dates(output_root, shop_id, biz_key)
            missing = compute_missing_dates(sz_start, sz_end, existing)
            print(f"\n>>> {biz_key} | 已存在 {len(existing)} 天 | 缺 {len(missing)} 天: {missing[:5]}{'...' if len(missing) > 5 else ''}")
            if not missing:
                print(f"[SKIP] {biz_key} 已覆盖，无需补齐")
                continue
            if dry_run:
                print(f"[DRY-RUN] {biz_key} | 需补 {len(missing)} 天")
                continue

            # 逐日循环补齐（商智服务端不支持多日区间导出，必须逐日调）
            for date_str in missing:
                print(f"\n -> {date_str}")
                try:
                    # ⚠️ 商智关键词分析需传 granularity=day（与近 30 天窗口对齐）
                    if biz_key == "商智关键词分析":
                        run_business(biz_key, date=date_str, granularity="day")
                    else:
                        run_business(biz_key, date=date_str)
                    print(f"  [OK] {biz_key} | {date_str}")
                except (CookieExpiredError, H5stExpiredError) as e:
                    print(f"  [AUTH_EXPIRED] {biz_key} | {date_str}: {e}")
                    has_auth_expired = True
                    break  # 鉴权过期 → 该业务后续日期不再尝试
                except Exception as e:
                    print(f"  [BIZ_FAIL] {biz_key} | {date_str}: {e}")
                    overall_exit = EXIT_BIZ_FAIL

    # 鉴权过期优先于业务失败返回
    if has_auth_expired:
        return EXIT_AUTH_EXPIRED
    return overall_exit


def main():
    """rpa_run.py 主入口"""
    parser = argparse.ArgumentParser(
        prog="rpa_run",
        description="多店铺数据导入手动触发入口（项目22，2026-08-17）",
    )
    parser.add_argument("--shop", default=None, help="指定单个店铺；不传则遍历全部 3 个店")
    parser.add_argument("--dry-run", action="store_true", help="仅预览计划，不实际下载/入库")
    parser.add_argument("--sz-days", type=int, default=30, help="商智业务日期窗口天数（默认30，验证时可设7）")
    parser.add_argument("--skip-jm", action="store_true", help="跳过京麦业务（避开频控窗口，只跑京准通+商智）")
    parser.add_argument("--only-jm", action="store_true", help="只跑京麦业务（与 --skip-jm 互斥；跳过京准通+商智）")
    parser.add_argument("--shop-interval", type=int, default=None,
                        help="店间安全间隔秒数（避免 IP 风控）。"
                             "默认从 config.xlsx「全局配置」shop_interval_seconds 读取（缺省 60 秒）。"
                             "传 0 = 不等待（不推荐）。")
    args = parser.parse_args()

    # 互斥校验
    if args.skip_jm and args.only_jm:
        print("[ERR] --skip-jm 与 --only-jm 互斥，不能同时使用")
        return EXIT_SYSTEM_ERROR

    # IMAP 鉴权过期拦截（启动时已检测；dry-run 模式放过，仅预览不实际执行）
    if _IMAP_BLOCK_ERROR is not None and not args.dry_run:
        print(f"[BLOCK] 拒绝执行：{_IMAP_BLOCK_ERROR}")
        return EXIT_SYSTEM_ERROR

    # 决定本次跑哪些店
    # Phase 2.3：店铺清单从 config.xlsx「店铺清单」sheet 读取（替代硬编码 KNOWN_SHOPS）
    shops = [args.shop] if args.shop else list_shop_ids()
    print(f"[INIT] 准备跑 {len(shops)} 个店: {shops}")
    print(f"[INIT] DRY-RUN 模式: {args.dry_run}")
    if args.only_jm:
        print(f"[INIT] 模式: 仅跑京麦业务（--only-jm）")

    if args.dry_run:
        print(f"\n{'='*70}\n[DRY-RUN 模式] 仅预览计划，不实际下载\n{'='*70}")

    # 累计所有店的退出码
    shop_results = {}
    all_known_shops = list_shop_ids(enabled_only=False)  # 含停用店铺，用于校验 --shop 入参
    # 计算店间安全间隔（2026-08-21 项目23 新增）
    # - 京东多店铺同 IP 频繁请求会触发风控（频控/限流/601/403）
    # - 默认 60 秒（每个店跑完所有业务后等够长再开始下一个店）
    # - 单店模式（--shop 指定）或 dry-run 模式不等待
    shop_interval = _read_shop_interval_seconds(args.shop_interval)
    is_single_shop = (len(shops) == 1)
    mode_parts = []
    if is_single_shop:
        mode_parts.append("单店模式，不等待")
    else:
        mode_parts.append(f"多店模式，每个店跑完后等待")
    if args.dry_run:
        mode_parts.append("dry-run 模式不等待")
    print(f"\n[INIT] 店间安全间隔：{shop_interval} 秒（{' / '.join(mode_parts)}）")

    for idx, shop_id in enumerate(shops):
        # 店间间隔（第 2 个店及之后，且非 dry-run，等待）
        if idx > 0 and not args.dry_run:
            print(f"\n⏸️  [COOLDOWN] 等待 {shop_interval} 秒后开始下一个店（避免 IP 风控）...")
            for remaining in range(shop_interval, 0, -10):
                if remaining % 30 == 0 or remaining <= 10:
                    print(f"   剩余 {remaining} 秒...")
                time.sleep(min(10, remaining))
            print(f"✅ [COOLDOWN] 冷却完成，开始跑 {shop_id}\n")

        if shop_id not in all_known_shops:
            print(f"[WARN] 未知店铺 {shop_id}（不在 config.xlsx「店铺清单」sheet 中），跳过")
            shop_results[shop_id] = EXIT_SYSTEM_ERROR
            continue
        try:
            shop_results[shop_id] = run_shop(shop_id, dry_run=args.dry_run, sz_days=args.sz_days, skip_jm=args.skip_jm, only_jm=args.only_jm)
        except Exception as e:
            print(f"[SYSTEM_ERROR] {shop_id}: {e}")
            traceback.print_exc()
            shop_results[shop_id] = EXIT_SYSTEM_ERROR

    # 输出汇总
    print(f"\n{'='*70}")
    print(f"[SUMMARY] {len(shops)} 个店执行结果")
    print(f"{'='*70}")
    for shop_id, code in shop_results.items():
        tag = {
            0: "✅ 成功",
            1: "⚠️ 部分失败",
            2: "🔄 鉴权过期（需RPA重抓）",
            3: "❌ 系统错误",
        }.get(code, f"?({code})")
        print(f"  {shop_id}: {tag}")

    # 决定最终退出码（任一店鉴权过期 → 整体返 2，触发影刀重抓）
    if any(c == EXIT_AUTH_EXPIRED for c in shop_results.values()):
        return EXIT_AUTH_EXPIRED
    if any(c == EXIT_SYSTEM_ERROR for c in shop_results.values()):
        return EXIT_SYSTEM_ERROR
    if any(c == EXIT_BIZ_FAIL for c in shop_results.values()):
        return EXIT_BIZ_FAIL
    return EXIT_SUCCESS


if __name__ == "__main__":
    sys.exit(main())