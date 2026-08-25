# -*- coding: utf-8 -*-
"""
创建 config.xlsx 配置文件

A列项目名 = 变量归属：
    全局：所有API共用的（路径/请求控制）
    商品搜索效果：只在本项目API中使用的（盐值/日期）

后续新增其他项目（如首页流量）时，可以再加新的项目名分类
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

output_dir = os.path.join(os.path.dirname(__file__), "config")
os.makedirs(output_dir, exist_ok=True)
config_path = os.path.join(output_dir, "config.xlsx")

wb = Workbook()
ws = wb.active
ws.title = "全局配置"

ws_headers = ["项目名", "变量参数", "参数值", "说明"]

ws_data = [
    # ===== 全局变量（所有API共用） =====
    ["全局", "cookie文件路径", "config/cookie.txt", "Cookie存放文件路径（相对项目根目录）"],
    ["全局", "输出目录", "output/", "导出文件保存路径"],
    ["全局", "日志目录", "logs/", "日志文件保存路径"],
    ["全局", "请求间隔(秒)", "30", "两次API调用之间的最小间隔"],
    ["全局", "最大重试次数", "3", "请求失败时的最大重试次数"],
    ["全局", "请求超时(秒)", "30", "单次请求的超时时间"],

    # ===== 商品搜索效果项目（仅本项目使用） =====
    ["商品搜索效果", "签名盐值", "372ad2c2b6", "京东MD5签名盐值。本项目使用，京东更新后需重新从commons.js逆向获得新值"],
    ["商品搜索效果", "date", "2026-08-03", "查询日期，格式YYYY-MM-DD"],
    ["商品搜索效果", "startDate", "2026-08-03", "开始日期，单日查询时与date相同"],
    ["商品搜索效果", "endDate", "2026-08-03", "结束日期，单日查询时与date相同"],
]

header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
data_font = Font(name="微软雅黑", size=10)
data_align = Alignment(vertical="top", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
for col_idx, header in enumerate(ws_headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font; cell.fill = header_fill; cell.alignment = header_align; cell.border = thin_border
for row_idx, row_data in enumerate(ws_data, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = data_font; cell.alignment = data_align; cell.border = thin_border
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 60
ws.column_dimensions['D'].width = 50
ws.freeze_panes = "A2"

wb.save(config_path)
print(f"✅ 配置文件已创建: {config_path}")
print(f"   单Sheet: 全局配置 - {len(ws_data)} 项")
print(f"   全局: 6项  商品搜索效果: 4项")
wb.close()