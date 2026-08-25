# -*- coding: utf-8 -*-
"""excel_master 自测脚本（2026-08-22 项目24）

4 个核心场景（用户 spec 第九节）：
    1. full_30day：覆盖已有日期 + 新增日期追加
    2. 90 万阈值：触发清理最早 60 天
    3. 空报表：跳过 sheet 不动其他
    4. repair_missing：补录模式覆盖追加

每个场景独立运行，不依赖京东真实接口，全部用 pandas DataFrame mock 数据。

用法：
    python tests/excel_master_unittest.py
    python tests/excel_master_unittest.py --scenario 1   # 只跑场景 1
    python tests/excel_master_unittest.py --keep        # 保留测试文件（默认测完删除）

依赖：
    - config.xlsx 已有「Excel总表映射」sheet（项目 24 一次性创建）
    - 临时目录 output/总表/_test_<shop_pin>/（自动创建，测完默认删除）
"""
import os
import sys
import shutil
import tempfile
import argparse
import logging
from datetime import datetime, timedelta

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

# 必须在 import excel_master 之前设置 SHOP_PIN（H-03 多店隔离）
TEST_SHOP_PIN = "TEST_PIN_8888"
TEST_SHOP_ID = "测试旗舰店"
os.environ["SHOP_ID"] = TEST_SHOP_ID
os.environ["SHOP_PIN"] = TEST_SHOP_PIN

import pandas as pd
import openpyxl

import excel_master as em


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)


def _make_test_dataframe(report_date: str, n_rows: int = 3) -> pd.DataFrame:
    """生成测试用 DataFrame。

    列名匹配 config.xlsx「Excel总表映射」中
    「商品流量来源_搜索」→ sheet「搜索流量」column_order 的格式。
    """
    return pd.DataFrame({
        "stat_date": [report_date] * n_rows,
        "商品名称": [f"测试商品{i+1}" for i in range(n_rows)],
        "商品SKU": [f"SKU{i+1:05d}" for i in range(n_rows)],
        "访客数": [100 * (i + 1) for i in range(n_rows)],
        "浏览量": [500 * (i + 1) for i in range(n_rows)],
        "平均停留时长": [60 + i for i in range(n_rows)],
        "加购客户数": [10 * (i + 1) for i in range(n_rows)],
        "加购转化率": [0.05 * (i + 1) for i in range(n_rows)],
        "成交客户数": [5 * (i + 1) for i in range(n_rows)],
        "成交转化率": [0.02 * (i + 1) for i in range(n_rows)],
        "成交金额": [1000 * (i + 1) for i in range(n_rows)],
        "客单价": [200 * (i + 1) for i in range(n_rows)],
        "下单客户数": [3 * (i + 1) for i in range(n_rows)],
        "下单金额": [600 * (i + 1) for i in range(n_rows)],
        "店铺关注人数": [50 * (i + 1) for i in range(n_rows)],
    })


def _redirect_master_dir_to_temp():
    """临时把 MASTER_DIR 指向 _test_<shop_pin>/ 目录。

    excel_master. MASTER_DIR 在 import 时就固定了，无法改全局，所以测试完后清理。
    """
    test_dir = os.path.join(PROJECT_ROOT, "output", "总表", f"_test_{TEST_SHOP_PIN}")
    if os.path.exists(test_dir):
        shutil.rmtree(test_dir)
    os.makedirs(test_dir, exist_ok=True)
    em.MASTER_DIR = test_dir
    return test_dir


def _read_sheet_rows(shop_pin: str, file_type: str, sheet_name: str) -> int:
    """读 sheet 数据行数（不含表头）。"""
    file_path = em._get_file_path(shop_pin, file_type)
    if not os.path.exists(file_path):
        return 0
    wb = openpyxl.load_workbook(file_path, read_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return 0
        ws = wb[sheet_name]
        # 减 1 去掉表头
        return max(0, ws.max_row - 1)
    finally:
        wb.close()


def _print_pass(scenario: str, msg: str):
    print(f"  ✅ [{scenario}] {msg}")


def _print_fail(scenario: str, msg: str):
    print(f"  ❌ [{scenario}] {msg}")
    raise AssertionError(msg)


# ============================================================
#  场景 1：full_30day — 覆盖已有日期 + 新增日期追加
# ============================================================
def scenario_full_30day() -> bool:
    """场景 1：full_30day 模式 — 先写 5 天，再写 3 天（其中 2 天重叠），验证覆盖+追加。"""
    scenario = "1-full_30day"
    print(f"\n{'='*70}\n🧪 [{scenario}] full_30day 覆盖+追加\n{'='*70}")

    test_dir = _redirect_master_dir_to_temp()
    em.reset_collected(TEST_SHOP_PIN)

    base_date = datetime(2026, 7, 1)
    # 第一波：写 5 天
    for i in range(5):
        d = (base_date + timedelta(days=i)).strftime("%Y-%m-%d")
        em.collect("商品流量来源_搜索", _make_test_dataframe(d, n_rows=3), d)
    _results = em.flush_shop(TEST_SHOP_ID, TEST_SHOP_PIN)
    n1 = _read_sheet_rows(TEST_SHOP_PIN, "sz", "搜索流量")
    _print_pass(scenario, f"第一波写入：sheet「搜索流量」共 {n1} 行")
    assert n1 == 15, f"期望 15 行（5天×3行），实际 {n1}"

    # 第二波：写 3 天（其中 7-03、7-04 重复；7-07 新增）
    em.reset_collected(TEST_SHOP_PIN)
    for d in ["2026-07-03", "2026-07-04", "2026-07-07"]:
        em.collect("商品流量来源_搜索", _make_test_dataframe(d, n_rows=4), d)
    em.flush_shop(TEST_SHOP_ID, TEST_SHOP_PIN)
    n2 = _read_sheet_rows(TEST_SHOP_PIN, "sz", "搜索流量")
    _print_pass(scenario, f"第二波写入后：共 {n2} 行")
    # 期望 = 7-01(3) + 7-02(3) + 7-03(4) + 7-04(4) + 7-05(3) + 7-07(4) = 21 行
    assert n2 == 21, f"期望 21 行（7-01+7-02+7-03覆盖+7-04覆盖+7-05+7-07新增），实际 {n2}"

    # 验证日期覆盖（7-03 应该是 4 行，不是 3+4=7）
    wb = openpyxl.load_workbook(em._get_file_path(TEST_SHOP_PIN, "sz"), read_only=True)
    try:
        ws = wb["搜索流量"]
        # 找 stat_date 列
        header = [c.value for c in ws[1]]
        sidx = header.index("stat_date") + 1
        from collections import Counter
        col_letter = openpyxl.utils.get_column_letter(sidx)
        date_count = Counter()
        for r in range(2, ws.max_row + 1):
            v = ws[f"{col_letter}{r}"].value
            if v:
                d = v.strftime("%Y-%m-%d") if isinstance(v, datetime) else str(v).strip()
                date_count[d] += 1
    finally:
        wb.close()
    print(f"     日期分布：{dict(sorted(date_count.items()))}")
    assert date_count["2026-07-03"] == 4, f"7-03 应是 4 行（覆盖），实际 {date_count['2026-07-03']}"
    assert date_count["2026-07-07"] == 4, f"7-07 应是 4 行（新增），实际 {date_count['2026-07-07']}"

    _print_pass(scenario, "日期覆盖+新增逻辑全部正确")
    return True


# ============================================================
#  场景 2：90 万阈值清理
# ============================================================
def scenario_threshold_cleanup() -> bool:
    """场景 2：模拟 90 万行触发，按 stat_date 删最早 60 天。

    性能优化：直接临时把 ROW_THRESHOLD 调到小值（比如 100），造 200 行数据触发清理。
    测试完恢复原阈值 900000。
    验证：
        1. 行数 >= ROW_THRESHOLD 时触发清理
        2. 按 stat_date 升序取第 60 个日期为 cutoff
        3. cutoff 及之前的所有行被删除
        4. cutoff 之后的行保留
    """
    scenario = "2-threshold_cleanup"
    print(f"\n{'='*70}\n🧪 [{scenario}] 90 万阈值清理（临时降低阈值加速测试）\n{'='*70}")

    test_dir = _redirect_master_dir_to_temp()
    em.reset_collected(TEST_SHOP_PIN)

    # 1. 临时把阈值调小（100 行）+ 清理窗口 5 天，加速测试
    original_threshold = em.ROW_THRESHOLD
    original_cleanup = em.CLEANUP_DAYS
    em.ROW_THRESHOLD = 100
    em.CLEANUP_DAYS = 5
    try:
        # 2. 造 200 行数据（100 天 × 2 行/天），覆盖第 1 天到第 100 天
        file_path = em._get_file_path(TEST_SHOP_PIN, "sz")
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        ws = wb.create_sheet("搜索流量")
        ws.append(["stat_date", "商品名称", "商品SKU"])

        n_rows_per_day = 2
        n_days = 100
        base_date = datetime(2026, 1, 1)
        for day_offset in range(n_days):
            d = (base_date + timedelta(days=day_offset)).strftime("%Y-%m-%d")
            for r in range(n_rows_per_day):
                ws.append([d, f"商品{day_offset}_{r}", f"SKU{r:05d}"])

        wb.save(file_path)
        wb.close()
        _print_pass(scenario, f"已造 {n_rows_per_day * n_days} 行测试数据（{n_days} 天 × {n_rows_per_day} 行，阈值 100）")

        # 3. 触发 flush → 应自动清理最早 5 天
        em.collect("商品流量来源_搜索", _make_test_dataframe("2026-04-15", n_rows=10), "2026-04-15")
        results = em.flush_shop(TEST_SHOP_ID, TEST_SHOP_PIN)

        sheet_result = results.get("sz", {}).get("搜索流量", {})
        cleanup_deleted = sheet_result.get("cleanup_deleted", 0)
        _print_pass(scenario, f"阈值清理：删除 {cleanup_deleted} 行")
        # 期望：清理最早 5 天 × 2 行 = 10 行
        assert cleanup_deleted == em.CLEANUP_DAYS * n_rows_per_day, (
            f"期望清理 {em.CLEANUP_DAYS * n_rows_per_day} 行，实际 {cleanup_deleted}"
        )

        # 4. 验证剩余行数
        n_after = _read_sheet_rows(TEST_SHOP_PIN, "sz", "搜索流量")
        # 期望：原 200 - 10 + 10（新 collect）= 200 行
        expected = (n_days - em.CLEANUP_DAYS) * n_rows_per_day + 10
        _print_pass(scenario, f"清理后剩余：{n_after} 行（期望 {expected}）")
        assert n_after == expected, f"期望 {expected} 行，实际 {n_after}"

        # 5. 验证清理后最早日期 >= 第 6 天（即 cutoff 之后的行保留）
        wb2 = openpyxl.load_workbook(file_path, read_only=True)
        try:
            ws2 = wb2["搜索流量"]
            from collections import Counter
            date_counter = Counter()
            header = [c.value for c in ws2[1]]
            sidx = header.index("stat_date") + 1
            from openpyxl.utils import get_column_letter
            cl = get_column_letter(sidx)
            for r in range(2, ws2.max_row + 1):
                v = ws2[f"{cl}{r}"].value
                if v:
                    d = v.strftime("%Y-%m-%d") if isinstance(v, datetime) else str(v).strip()
                    date_counter[d] += 1
        finally:
            wb2.close()
        # 最早日期应该是第 6 天（2026-01-06）
        dates_sorted = sorted(date_counter.keys())
        assert dates_sorted[0] == "2026-01-06", f"最早日期期望 2026-01-06，实际 {dates_sorted[0]}"
        _print_pass(scenario, f"清理后最早日期：{dates_sorted[0]}（符合预期）")

    finally:
        # 6. 恢复原阈值（防止污染其他测试/真实业务）
        em.ROW_THRESHOLD = original_threshold
        em.CLEANUP_DAYS = original_cleanup
        _print_pass(scenario, f"已恢复 ROW_THRESHOLD={original_threshold}, CLEANUP_DAYS={original_cleanup}")

    return True


# ============================================================
#  场景 3：空报表 — 跳过 sheet 不动其他
# ============================================================
def scenario_empty_data() -> bool:
    """场景 3：模拟空 df + 不影响其他 sheet。"""
    scenario = "3-empty_data"
    print(f"\n{'='*70}\n🧪 [{scenario}] 空报表跳过\n{'='*70}")

    test_dir = _redirect_master_dir_to_temp()
    em.reset_collected(TEST_SHOP_PIN)

    # 第一次：写搜索流量 + 推荐流量 2 个 sheet
    em.collect("商品流量来源_搜索", _make_test_dataframe("2026-08-01", n_rows=3), "2026-08-01")
    em.collect("商品流量来源_推荐", _make_test_dataframe("2026-08-01", n_rows=3), "2026-08-01")
    em.flush_shop(TEST_SHOP_ID, TEST_SHOP_PIN)

    n_search = _read_sheet_rows(TEST_SHOP_PIN, "sz", "搜索流量")
    n_recommend = _read_sheet_rows(TEST_SHOP_PIN, "sz", "推荐流量")
    _print_pass(scenario, f"初始状态：搜索={n_search}, 推荐={n_recommend}")
    assert n_search == 3 and n_recommend == 3

    # 第二次：只 collect 搜索流量（推荐流量不 collect = 空数据 = 跳过）
    em.reset_collected(TEST_SHOP_PIN)
    em.collect("商品流量来源_搜索", _make_test_dataframe("2026-08-02", n_rows=3), "2026-08-02")
    # 故意调 collect 空 df（业务类有时会传空 df 进来）
    em.collect("商品流量来源_购物车", pd.DataFrame(), "2026-08-02")  # 空 df
    results = em.flush_shop(TEST_SHOP_ID, TEST_SHOP_PIN)

    sheet_result = results.get("sz", {}).get("购物车流量", {})
    _print_pass(scenario, f"购物车流量 sheet 状态：{sheet_result}")
    # 期望：购物车流量 sheet 不被创建（skipped=True）
    assert sheet_result.get("skipped") is True, f"购物车流量应被跳过，实际 {sheet_result}"

    # 验证搜索流量：8-01(3) + 8-02(3) = 6 行
    n_search2 = _read_sheet_rows(TEST_SHOP_PIN, "sz", "搜索流量")
    _print_pass(scenario, f"搜索流量更新后：{n_search2} 行")
    assert n_search2 == 6, f"期望 6 行（覆盖+新增），实际 {n_search2}"

    # 验证推荐流量：8-01(3) 行不变
    n_recommend2 = _read_sheet_rows(TEST_SHOP_PIN, "sz", "推荐流量")
    _print_pass(scenario, f"推荐流量未被影响：{n_recommend2} 行")
    assert n_recommend2 == 3, f"期望 3 行（未变），实际 {n_recommend2}"

    return True


# ============================================================
#  场景 4：repair_missing — 补录模式覆盖追加
# ============================================================
def scenario_repair_missing() -> bool:
    """场景 4：补录模式 — 已有数据基础上补几天，验证只追加缺失日期。"""
    scenario = "4-repair_missing"
    print(f"\n{'='*70}\n🧪 [{scenario}] repair_missing 补录\n{'='*70}")

    test_dir = _redirect_master_dir_to_temp()
    em.reset_collected(TEST_SHOP_PIN)

    # 先写 10 天数据
    for i in range(10):
        d = (datetime(2026, 6, 1) + timedelta(days=i)).strftime("%Y-%m-%d")
        em.collect("店铺来源_三级渠道", _make_three_channel_dataframe(d, n_rows=2), d)
    em.flush_shop(TEST_SHOP_ID, TEST_SHOP_PIN)

    n1 = _read_sheet_rows(TEST_SHOP_PIN, "sz", "店铺来源_三级渠道")
    _print_pass(scenario, f"初次写入：{n1} 行（10天×2行）")
    assert n1 == 20

    # 补录：6-05 已存在（应被覆盖），6-15 新增，6-20 新增
    em.reset_collected(TEST_SHOP_PIN)
    for d in ["2026-06-05", "2026-06-15", "2026-06-20"]:
        em.collect("店铺来源_三级渠道", _make_three_channel_dataframe(d, n_rows=3), d)
    em.flush_shop(TEST_SHOP_ID, TEST_SHOP_PIN)

    n2 = _read_sheet_rows(TEST_SHOP_PIN, "sz", "店铺来源_三级渠道")
    _print_pass(scenario, f"补录后：{n2} 行")
    # 期望：原 20 + 6-05 新 3 替换原 2 = 21; 6-15(+3), 6-20(+3) → 27
    assert n2 == 27, f"期望 27 行，实际 {n2}"

    return True


def _make_three_channel_dataframe(report_date: str, n_rows: int = 3) -> pd.DataFrame:
    """三级渠道测试用 df。"""
    return pd.DataFrame({
        "stat_date": [report_date] * n_rows,
        "三级渠道名称": [f"渠道{i+1}" for i in range(n_rows)],
        "三级渠道ID": [1000 + i for i in range(n_rows)],
        "进店访客数": [100 * (i + 1) for i in range(n_rows)],
        "进店浏览量": [500 * (i + 1) for i in range(n_rows)],
        "访客数占比": [0.1 * (i + 1) for i in range(n_rows)],
        "浏览量占比": [0.05 * (i + 1) for i in range(n_rows)],
        "成交客户数": [10 * (i + 1) for i in range(n_rows)],
        "成交金额": [1000 * (i + 1) for i in range(n_rows)],
        "成交转化率": [0.02 * (i + 1) for i in range(n_rows)],
    })


# ============================================================
#  主入口
# ============================================================
SCENARIOS = {
    "1": ("full_30day 覆盖+追加", scenario_full_30day),
    "2": ("90 万阈值清理", scenario_threshold_cleanup),
    "3": ("空报表跳过", scenario_empty_data),
    "4": ("repair_missing 补录", scenario_repair_missing),
}


# 阈值清理测试期望参数（项目24 excel_master.ROW_THRESHOLD=900000 CLEANUP_DAYS=60）
CLEANUP_DAYS_EXPECTED = 60


def main():
    parser = argparse.ArgumentParser(description="excel_master 自测脚本")
    parser.add_argument("--scenario", choices=list(SCENARIOS.keys()), default=None,
                        help="只跑指定场景（默认全部跑）")
    parser.add_argument("--keep", action="store_true",
                        help="保留测试文件（默认测完删除 output/总表/_test_*/）")
    args = parser.parse_args()

    print("=" * 70)
    print("🧪 Excel 总表自测 (项目24，2026-08-22)")
    print(f"   测试店铺：{TEST_SHOP_ID} ({TEST_SHOP_PIN})")
    print(f"   临时目录：output/总表/_test_{TEST_SHOP_PIN}/")
    print("=" * 70)

    # 校验 config.xlsx 已有「Excel总表映射」sheet
    try:
        mappings = em._load_mapping_from_xlsx()
        print(f"   映射配置：{len(mappings)} 条启用")
        from collections import Counter
        ft_count = Counter(m["file_type"] for m in mappings)
        for ft, n in sorted(ft_count.items()):
            print(f"     - {ft}: {n} sheets")
    except Exception as e:
        print(f"❌ 配置读取失败：{e}")
        print("   → 请先运行 _add_excel_master_mapping.py（一次性脚本）")
        sys.exit(1)

    success = 0
    failed = 0
    to_run = [args.scenario] if args.scenario else list(SCENARIOS.keys())
    for s in to_run:
        desc, fn = SCENARIOS[s]
        try:
            fn()
            success += 1
        except AssertionError as e:
            print(f"  ❌ [{s}] 断言失败：{e}")
            failed += 1
        except Exception as e:
            import traceback
            traceback.print_exc()
            print(f"  ❌ [{s}] 异常：{type(e).__name__}: {e}")
            failed += 1

    # 清理
    if not args.keep:
        test_dir = os.path.join(PROJECT_ROOT, "output", "总表", f"_test_{TEST_SHOP_PIN}")
        if os.path.exists(test_dir):
            shutil.rmtree(test_dir)
            print(f"\n🧹 清理测试目录：{test_dir}")
    else:
        print(f"\n💾 保留测试目录：output/总表/_test_{TEST_SHOP_PIN}/")

    # 汇总
    print("\n" + "=" * 70)
    total = len(to_run)
    print(f"📊 自测结果：✅ {success} / ❌ {failed} / 总 {total}")
    print("=" * 70)
    sys.exit(0 if failed == 0 else 1)


if __name__ == "__main__":
    main()