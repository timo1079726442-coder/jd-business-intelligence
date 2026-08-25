# -*- coding: utf-8 -*-
"""biz_config_loader.py
============================================================
店铺/业务清单统一加载器（2026-08-20 Phase 2.2 上线）

目的：
    - 替代 rpa_run.py / daily_update.py / fill_missing.py 中硬编码的
      KNOWN_SHOPS / JM_JZT_BIZ_KEYS / 商智_BIZ_KEYS / biz_to_dirname /
      MVP_BIZ_KEYS / BIZ_FEATURES
    - 替代 auth_loader.py 的 SHOP_ID_TO_PREFIX
    - 单一数据源：config.xlsx「店铺清单」+「业务清单」两个 sheet
    - 兜底机制：旧版 config.xlsx 没有这两个 sheet 时，回退到代码兜底常量
      并打印警告提醒补写 config（AGENTS.md 第26条）

使用示例：
    from biz_config_loader import (
        list_shops, get_shop_prefix,
        list_biz_keys, get_biz_config, get_output_dirname,
    )

    # 获取所有启用的店铺
    shops = list_shops(enabled_only=True)   # → [ShopConfig(shop_id='FYA箱包旗舰店', ...)]

    # 店名 → RPA 文件命名前缀
    prefix = get_shop_prefix("FYA箱包旗舰店")  # → "{{FYA}}"

    # 按 batch 分组的业务清单
    jm_jzt_keys = list_biz_keys(batch="jm_jzt", enabled_only=True)
    sz_keys = list_biz_keys(batch="sz", enabled_only=True)

    # 业务 → 输出目录名（替代 rpa_run.py 的 biz_to_dirname）
    dirname = get_output_dirname("京麦订单明细_完整一键导出")  # → "京麦订单明细"

AGENTS.md 合规：
    - 第3条：业务配置从 config.xlsx「全局配置」sheet 读取，通过 biz_key 匹配
    - 第26条：业务参数优先 config.xlsx；代码兜底值必须打印警告
    - 第50条 skill 规则：禁止在调度脚本中硬编码业务列表
"""
import os
import logging
from dataclasses import dataclass
from typing import List, Optional

# ====================================================================
#  数据类定义
# ====================================================================

@dataclass
class ShopConfig:
    """店铺配置（对应 config.xlsx「店铺清单」sheet 一行）"""
    shop_id: str      # 店铺全名，如 "FYA箱包旗舰店"
    file_prefix: str  # RPA 文件命名前缀，如 "{{FYA}}"
    enabled: bool     # 是否启用
    jzt_pin_options: str = ""  # 京准通子账号列表（逗号分隔"账号名:账号ID"），如 "FYA8888:99936530475,..."

    @property
    def short_name(self) -> str:
        """店铺短名（去掉"箱包旗舰店"后缀），如 "FYA" """
        return self.shop_id.replace("箱包旗舰店", "")


@dataclass
class BizConfig:
    """业务配置（对应 config.xlsx「业务清单」sheet 一行）"""
    biz_key: str             # 业务 key，如 "京麦订单明细_完整一键导出"
    batch: str               # 调度批次：jm_jzt=近30天全量 / sz=缺日补齐
    output_dirname: str      # 输出目录名，如 "京麦订单明细"
    supports_range: bool     # 是否支持区间调用（False=逐日循环）
    default_granularity: Optional[str]  # 默认粒度（day/month/None）
    enabled: bool            # 是否启用


# ====================================================================
#  代码兜底常量（AGENTS.md 第26条 兜底值，缺 config 时回退 + 警告）
# ====================================================================
# ⚠️ 警告：以下兜底数据仅在 config.xlsx 缺少「店铺清单」「业务清单」sheet 时使用
# 请勿直接修改这里的常量来增删店铺/业务，应修改 config.xlsx 后让本模块自动读取

_FALLBACK_SHOPS = [
    # shop_id, file_prefix, enabled, jzt_pin_options
    ("FYA箱包旗舰店",  "{{FYA}}",  True,
     "FYA8888:99936530475,FYA888888:99936525688,FAY掌柜888:99937142699,FYA19529975351:99938531397,fya掌柜777:99938957251,FYA小婷:99938963919,FYA少冰:99945916633,FYA小冠:99947097388,FYA布丁:99952884963,FYA小柔:99955522890,FYA小敏:99960125485"),
    # MIYO 店主账号（用户提供，2026-08-20）：pin=miyo-周, 数字ID=99918672844
    ("MIYO箱包旗舰店", "{{MIYO}}", True, "miyo-周:99918672844"),
    # OTA 店主账号（用户提供，2026-08-20）：pin=ota8888, 数字ID=99911675921
    ("OTA箱包旗舰店",  "{{OTA}}",  True, "ota8888:99911675921"),
]

_FALLBACK_BIZ = [
    # biz_key, batch, output_dirname, supports_range, default_granularity, enabled
    ("京麦订单明细_完整一键导出",      "jm_jzt", "京麦订单明细",           True,  None,  True),
    ("京麦售后明细_完整一键导出",      "jm_jzt", "京麦售后明细",           True,  None,  True),
    ("京准通快车自定义报表",           "jm_jzt", "京准通快车效果自定义",   True,  None,  True),
    ("京准通快车订单效果明细",         "jm_jzt", "京准通快车订单效果明细", True,  None,  True),
    ("京准通全站营销单品计划",         "jm_jzt", "京准通全站营销单品计划", True,  None,  True),
    ("京准通全站营销单品推广效果",     "jm_jzt", "京准通全站营销单品推广效果", True, None, True),
    ("京准通全站营销全店计划",         "jm_jzt", "京准通全站营销全店计划", True,  None,  True),
    ("京准通全站营销全店推广效果",     "jm_jzt", "京准通全站营销全店推广效果", True, None, True),
    ("商品流量来源_搜索",              "sz",     "搜索流量",              False, None,  True),
    ("商品流量来源_购物车",            "sz",     "购物车流量",            False, None,  True),
    ("商品流量来源_推荐",              "sz",     "推荐流量",              False, None,  True),
    ("商品流量来源_自主访问",          "sz",     "自主访问流量",          False, None,  False),  # 停用业务
    ("店铺来源_三级渠道",              "sz",     "店铺来源_三级渠道",     True,  None,  True),
    ("商品明细导出",                   "sz",     "商品明细",              True,  None,  True),
    ("商品流失分析",                   "sz",     "商品流失分析",          True,  None,  True),
    ("商智关键词分析",                 "sz",     "商智关键词分析",        False, "day", True),  # day 粒度跑日表
]


# ====================================================================
#  config.xlsx 读取
# ====================================================================

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(PROJECT_ROOT, "config", "config.xlsx")


def _str_to_bool(val) -> bool:
    """字符串转 bool（兼容 是/否、True/False、1/0）"""
    if isinstance(val, bool):
        return val
    s = str(val).strip().lower()
    return s in ("是", "true", "1", "yes", "y", "on")


def _load_shops_from_xlsx() -> Optional[List[ShopConfig]]:
    """从 config.xlsx「店铺清单」sheet 读取店铺列表。

    Returns:
        list[ShopConfig] - 成功读取的店铺列表
        None             - sheet 不存在或读取失败（调用方应回退兜底）
    """
    if not os.path.exists(CONFIG_PATH):
        return None
    try:
        import openpyxl
        wb = openpyxl.load_workbook(CONFIG_PATH, read_only=True, data_only=True)
        if "店铺清单" not in wb.sheetnames:
            wb.close()
            return None
        ws = wb["店铺清单"]
        shops = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:  # 跳过表头
                continue
            if not row or len(row) < 3 or not row[0]:
                continue
            # 第 4 列 jzt_pin_options（逗号分隔"账号名:账号ID"），可空
            jzt_pin = str(row[3]).strip() if len(row) >= 4 and row[3] else ""
            shops.append(ShopConfig(
                shop_id=str(row[0]).strip(),
                file_prefix=str(row[1]).strip() if row[1] else "",
                enabled=_str_to_bool(row[2]),
                jzt_pin_options=jzt_pin,
            ))
        wb.close()
        return shops if shops else None
    except Exception as e:
        logging.warning(f"⚠️ config.xlsx「店铺清单」sheet 读取失败：{e}")
        return None


def _load_biz_from_xlsx() -> Optional[List[BizConfig]]:
    """从 config.xlsx「业务清单」sheet 读取业务列表。

    Returns:
        list[BizConfig] - 成功读取的业务列表
        None            - sheet 不存在或读取失败
    """
    if not os.path.exists(CONFIG_PATH):
        return None
    try:
        import openpyxl
        wb = openpyxl.load_workbook(CONFIG_PATH, read_only=True, data_only=True)
        if "业务清单" not in wb.sheetnames:
            wb.close()
            return None
        ws = wb["业务清单"]
        biz_list = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue
            if not row or len(row) < 6 or not row[0]:
                continue
            granularity = str(row[4]).strip() if row[4] else None
            granularity = granularity if granularity else None
            biz_list.append(BizConfig(
                biz_key=str(row[0]).strip(),
                batch=str(row[1]).strip() if row[1] else "sz",
                output_dirname=str(row[2]).strip() if row[2] else str(row[0]).strip(),
                supports_range=_str_to_bool(row[3]),
                default_granularity=granularity,
                enabled=_str_to_bool(row[5]),
            ))
        wb.close()
        return biz_list if biz_list else None
    except Exception as e:
        logging.warning(f"⚠️ config.xlsx「业务清单」sheet 读取失败：{e}")
        return None


def _fallback_shops() -> List[ShopConfig]:
    """兜底店铺列表（config.xlsx 缺 sheet 时使用）"""
    print(
        "[WARN] ⚠️ config.xlsx 缺少「店铺清单」sheet，使用代码兜底常量。\n"
        "       ⚠️ 严禁长期依赖兜底！请运行 `python create_config_sheets.py` 生成 sheet。"
    )
    return [ShopConfig(s, p, e, pin) for s, p, e, pin in _FALLBACK_SHOPS]


def _fallback_biz() -> List[BizConfig]:
    """兜底业务列表"""
    print(
        "[WARN] ⚠️ config.xlsx 缺少「业务清单」sheet，使用代码兜底常量。\n"
        "       ⚠️ 严禁长期依赖兜底！请运行 `python create_config_sheets.py` 生成 sheet。"
    )
    return [BizConfig(k, b, d, r, g, e) for k, b, d, r, g, e in _FALLBACK_BIZ]


# ====================================================================
#  对外 API
# ====================================================================

def list_shops(enabled_only: bool = True) -> List[ShopConfig]:
    """返回店铺列表。

    Args:
        enabled_only: True=只返回 enabled=True 的店铺（默认）
                      False=返回所有店铺
    Returns:
        List[ShopConfig]
    """
    shops = _load_shops_from_xlsx() or _fallback_shops()
    if enabled_only:
        shops = [s for s in shops if s.enabled]
    return shops


def get_shop_prefix(shop_id: str) -> str:
    """shop_id → RPA 文件命名前缀。

    替代 auth_loader.SHOP_ID_TO_PREFIX 的查询。
    查不到时按 {{shop_id}} 兜底（与 auth_loader.resolve_file_prefix 一致）。

    示例:
        get_shop_prefix("FYA箱包旗舰店") → "{{FYA}}"
        get_shop_prefix("未知店") → "{{未知店}}"
    """
    shops = _load_shops_from_xlsx() or _fallback_shops()
    for s in shops:
        if s.shop_id == shop_id:
            return s.file_prefix
    # 兜底：未在 config 注册的店铺，按规则包双花括号
    return f"{{{{{shop_id}}}}}"


def list_shop_ids(enabled_only: bool = True) -> List[str]:
    """返回 shop_id 字符串列表（便捷方法，替代 rpa_run.py 的 KNOWN_SHOPS）"""
    return [s.shop_id for s in list_shops(enabled_only=enabled_only)]


def get_jzt_pin_options(shop_id: str) -> List[dict]:
    """解析 config.xlsx「店铺清单」jzt_pin_options 列，返回京准通 payload 需要的 options 列表。

    config.xlsx 格式：逗号分隔的"账号名:账号ID"，如 "FYA8888:99936530475,FYA888888:99936525688"
    返回格式（对应 JZT_KUAICHE_PAYLOAD_TEMPLATE.customDimensionOptions.pin.options[0].options）：
        [
            {"checked": True,  "desc": "FYA8888",   "flag": True,  "hidden": False, "key": "99936530475", "value": "FYA8888"},
            {"checked": False, "desc": "FYA888888", "flag": False, "hidden": False, "key": "99936525688", "value": "FYA888888"},
            ...
        ]
    规则：
        - 第一个账号 = 当前账号（flag=True, checked=True）
        - 其他账号 = flag=False, checked=False
        - config 为空 → 回退到兜底常量；兜底也为空 → 返回空列表（清空 pin.options，任务仍可创建）
    """
    # 1. 先从 config.xlsx 读取
    xlsx_shops = _load_shops_from_xlsx()
    raw = ""
    if xlsx_shops:
        for s in xlsx_shops:
            if s.shop_id == shop_id:
                raw = s.jzt_pin_options
                break
    # 2. config.xlsx 为空或未配置 → 回退到兜底常量
    if not raw:
        for s in _fallback_shops():
            if s.shop_id == shop_id:
                raw = s.jzt_pin_options
                break
    if not raw:
        return []
    options = []
    for i, item in enumerate(raw.split(",")):
        item = item.strip()
        if not item:
            continue
        # 解析"账号名:账号ID"；缺 ID 时 key 也用账号名（兜底）
        if ":" in item:
            desc, key = item.split(":", 1)
            desc, key = desc.strip(), key.strip()
        else:
            desc, key = item, item
        is_first = (i == 0)
        options.append({
            "checked": is_first,   # 第一个账号默认勾选（当前账号）
            "desc": desc,          # 账号名（前端显示）
            "flag": is_first,      # True=当前账号标记
            "hidden": False,
            "key": key,            # 账号ID（数字字符串）
            "value": desc,         # 账号名（与 desc 一致）
        })
    return options


def list_biz_keys(
    batch: Optional[str] = None,
    enabled_only: bool = True,
) -> List[str]:
    """返回 biz_key 字符串列表。

    Args:
        batch: 筛选批次（"jm_jzt" 或 "sz"）；None=所有批次
        enabled_only: True=只返回 enabled=True 的业务

    Returns:
        List[str] - biz_key 列表
    """
    biz_list = _load_biz_from_xlsx() or _fallback_biz()
    if enabled_only:
        biz_list = [b for b in biz_list if b.enabled]
    if batch is not None:
        biz_list = [b for b in biz_list if b.batch == batch]
    return [b.biz_key for b in biz_list]


def list_biz_configs(
    batch: Optional[str] = None,
    enabled_only: bool = True,
) -> List[BizConfig]:
    """返回 BizConfig 对象列表（带完整字段）"""
    biz_list = _load_biz_from_xlsx() or _fallback_biz()
    if enabled_only:
        biz_list = [b for b in biz_list if b.enabled]
    if batch is not None:
        biz_list = [b for b in biz_list if b.batch == batch]
    return biz_list


def get_biz_config(biz_key: str) -> Optional[BizConfig]:
    """单个业务配置查询。"""
    biz_list = _load_biz_from_xlsx() or _fallback_biz()
    for b in biz_list:
        if b.biz_key == biz_key:
            return b
    return None


def get_output_dirname(biz_key: str) -> str:
    """biz_key → 输出目录名（替代 rpa_run.py 的 biz_to_dirname）。

    查不到时返回 biz_key 本身（兜底）。
    """
    cfg = get_biz_config(biz_key)
    if cfg:
        return cfg.output_dirname
    return biz_key


def get_biz_feature(biz_key: str) -> dict:
    """获取业务特性字典（兼容 daily_update.py 的 BIZ_FEATURES 格式）。

    返回:
        {
            "supports_range": bool,
            "default_granularity": Optional[str],
            "extra_kwargs": {},
        }
        查不到返回空字典 {}，调用方应判断跳过。
    """
    cfg = get_biz_config(biz_key)
    if not cfg:
        return {}
    return {
        "supports_range": cfg.supports_range,
        "default_granularity": cfg.default_granularity,
        "extra_kwargs": {},
    }


# ====================================================================
#  CLI（调试用）
# ====================================================================
if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("用法：")
        print("  python biz_config_loader.py shops          # 列出所有店铺")
        print("  python biz_config_loader.py biz jm_jzt     # 列出 jm_jzt 批次业务")
        print("  python biz_config_loader.py biz sz         # 列出 sz 批次业务")
        print("  python biz_config_loader.py biz all        # 列出所有业务")
        print("  python biz_config_loader.py prefix <shop_id>  # 查店铺前缀")
        print("  python biz_config_loader.py dirname <biz_key>  # 查业务输出目录名")
        sys.exit(0)

    cmd = sys.argv[1]
    if cmd == "shops":
        print("=== 所有启用的店铺 ===")
        for s in list_shops(enabled_only=True):
            tag = "✅ 启用" if s.enabled else "❌ 停用"
            print(f"  {s.shop_id} | prefix={s.file_prefix} | {tag}")

    elif cmd == "biz":
        batch = sys.argv[2] if len(sys.argv) > 2 else "all"
        configs = list_biz_configs(
            batch=None if batch == "all" else batch,
            enabled_only=False,
        )
        print(f"=== 业务清单（batch={batch}）===")
        for b in configs:
            tag = "✅" if b.enabled else "❌"
            range_str = "区间" if b.supports_range else "逐日"
            print(f"  {tag} {b.biz_key:32s} | {b.batch} | {b.output_dirname:24s} | {range_str} | {b.default_granularity or '-'}")

    elif cmd == "prefix":
        if len(sys.argv) < 3:
            print("用法：python biz_config_loader.py prefix <shop_id>")
            sys.exit(1)
        print(f"{sys.argv[2]} → {get_shop_prefix(sys.argv[2])}")

    elif cmd == "dirname":
        if len(sys.argv) < 3:
            print("用法：python biz_config_loader.py dirname <biz_key>")
            sys.exit(1)
        print(f"{sys.argv[2]} → {get_output_dirname(sys.argv[2])}")

    else:
        print(f"未知命令：{cmd}")
